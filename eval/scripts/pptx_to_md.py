#!/usr/bin/env python3
"""
PPTX → Markdown 预处理（绕过 V3 PptxParser 的 NeedsMultimodalParsingException）。

V3 PptxParser 抽不到文本时会主动抛"需 OCR"异常；
本脚本用 python-pptx 直接抽每页文字 + 备注，作为替代路径。

用法：python3 pptx_to_md.py <input.pptx> <output.md>
"""
import sys
from pathlib import Path

from pptx import Presentation


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2])
    if not src.exists():
        print(f"❌ {src}")
        sys.exit(1)

    prs = Presentation(str(src))
    sections = [f"# {src.stem}", ""]
    for i, slide in enumerate(prs.slides, 1):
        title = ""
        body_lines = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                txt = "".join(run.text for run in para.runs).strip()
                if not txt:
                    continue
                # 拿首个非空段当 title
                if not title and len(txt) <= 80 and "：" not in txt:
                    title = txt
                else:
                    body_lines.append(txt)
        sections.append(f"## Slide {i}: {title or '(无标题)'}")
        sections.append("")
        for line in body_lines:
            sections.append(f"- {line}")
        # 备注
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            sections.append("")
            sections.append(f"> 备注: {slide.notes_slide.notes_text_frame.text.strip()}")
        sections.append("")

    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_text("\n".join(sections), encoding="utf-8")
    print(f"✅ {src.name} -> {dst.name} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
