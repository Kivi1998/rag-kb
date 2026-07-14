#!/usr/bin/env python3
"""
DOCX → Markdown 预处理（绕过 V3 DocxParser 对内嵌图片的拒答）。

用 python-docx 正确识别 heading（# ## ###）和样式，
保证 TechnicalChunker 切出的 chunk heading_path 不为空。

用法：python3 docx_to_md.py <input.docx> <output.md>
"""
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def heading_level(style_name: str) -> int:
    """从样式名解析 heading 级别（'Heading 1' → 1，'Heading 2' → 2 ...）。"""
    if not style_name:
        return 0
    if "Heading" in style_name or "标题" in style_name:
        for i in range(1, 7):
            if str(i) in style_name:
                return i
    return 0


def is_bold(run) -> bool:
    """判断 run 是否加粗（粗体可能是非标准 heading）。"""
    return bool(run.bold)


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2])
    if not src.exists():
        print(f"❌ {src}")
        sys.exit(1)

    doc = Document(str(src))
    md_lines = [f"# {src.stem}", ""]

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            md_lines.append("")
            continue
        # 优先用 paragraph.style 解析 heading
        style_name = para.style.name if para.style else ""
        level = heading_level(style_name)
        if level == 0 and para.runs and is_bold(para.runs[0]) and len(text) < 80:
            # 粗体短文本可能是非标准 heading
            level = 2
        if level > 0:
            md_lines.append("#" * level + " " + text)
            md_lines.append("")
        else:
            md_lines.append(text)
            md_lines.append("")

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                md_lines.append("| " + " | ".join(cells) + " |")
        md_lines.append("")

    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_text("\n".join(md_lines), encoding="utf-8")
    print(f"✅ {src.name} -> {dst.name} ({len(doc.paragraphs)} paragraphs)")


if __name__ == "__main__":
    main()
