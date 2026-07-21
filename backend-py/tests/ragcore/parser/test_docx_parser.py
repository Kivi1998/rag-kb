"""DocxParser 测试。

用 `python-docx` 现造最小 docx 文档验证解析结果结构（不是逐字节对齐 Java 版）：
标题层级追踪、表格提取、内嵌图片触发 `NeedsMultimodalParsingException`、无可索引
文本时同样触发该异常。文档现造后存到内存 `BytesIO` 再解析，不依赖任何现成语料文件。
"""

from __future__ import annotations

import io

import pytest
from docx import Document

from app.ragcore.model import NeedsMultimodalParsingException
from app.ragcore.parser.docx_parser import DocxParser
from tests.ragcore.parser.conftest import FakeDocumentInput

_DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _build_docx(build_fn) -> bytes:
    document = Document()
    build_fn(document)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _docx_input(data: bytes, filename: str = "doc.docx") -> FakeDocumentInput:
    return FakeDocumentInput(data, filename, content_type=_DOCX_CONTENT_TYPE)


class TestSupports:
    def test_supports_by_content_type(self):
        assert DocxParser().supports(_DOCX_CONTENT_TYPE, None) is True

    def test_supports_by_extension(self):
        assert DocxParser().supports(None, "report.docx") is True

    def test_rejects_unrelated(self):
        assert DocxParser().supports(None, "report.pdf") is False


class TestHeadingHierarchy:
    def test_two_paragraphs_under_same_heading(self):
        def build(doc):
            doc.add_heading("第一章 概述", level=1)
            doc.add_paragraph("这是第一段正文。")
            doc.add_paragraph("这是第二段正文。")

        data = _build_docx(build)
        result = DocxParser().parse(_docx_input(data))

        body_sections = [s for s in result.sections if "第一段正文" in s.text]
        assert len(body_sections) == 1
        assert body_sections[0].heading_path == ("第一章 概述",)
        assert "第二段正文" in body_sections[0].text

    def test_nested_heading_levels_build_path(self):
        def build(doc):
            doc.add_heading("第一章", level=1)
            doc.add_heading("1.1 小节", level=2)
            doc.add_paragraph("小节正文。")

        data = _build_docx(build)
        result = DocxParser().parse(_docx_input(data))

        body_section = next(s for s in result.sections if "小节正文" in s.text)
        assert body_section.heading_path == ("第一章", "1.1 小节")

    def test_sibling_heading_replaces_previous_at_same_level(self):
        def build(doc):
            doc.add_heading("第一章", level=1)
            doc.add_paragraph("第一章内容。")
            doc.add_heading("第二章", level=1)
            doc.add_paragraph("第二章内容。")

        data = _build_docx(build)
        result = DocxParser().parse(_docx_input(data))

        second_section = next(s for s in result.sections if "第二章内容" in s.text)
        assert second_section.heading_path == ("第二章",)
        assert "第一章" not in second_section.heading_path


class TestTableExtraction:
    def test_table_cells_joined_with_pipe(self):
        def build(doc):
            doc.add_heading("数据表", level=1)
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "列A"
            table.cell(0, 1).text = "列B"
            table.cell(1, 0).text = "1"
            table.cell(1, 1).text = "2"

        data = _build_docx(build)
        result = DocxParser().parse(_docx_input(data))

        table_sections = [s for s in result.sections if s.metadata.get("type") == "table"]
        assert len(table_sections) == 1
        assert table_sections[0].text == "列A | 列B\n1 | 2"
        assert table_sections[0].heading_path == ("数据表",)


class TestNoIndexableTextRaises:
    def test_empty_document_raises_needs_multimodal_parsing(self):
        data = _build_docx(lambda doc: None)
        with pytest.raises(NeedsMultimodalParsingException):
            DocxParser().parse(_docx_input(data))

    def test_whitespace_only_paragraphs_raise_needs_multimodal_parsing(self):
        def build(doc):
            doc.add_paragraph("   ")
            doc.add_paragraph("")

        data = _build_docx(build)
        with pytest.raises(NeedsMultimodalParsingException):
            DocxParser().parse(_docx_input(data))


class TestEmbeddedPictureRaises:
    def test_document_with_inline_picture_raises_needs_multimodal_parsing(self, tmp_path):
        from PIL import Image

        image_path = tmp_path / "figure.png"
        Image.new("RGB", (4, 4), color="red").save(image_path)

        def build(doc):
            doc.add_paragraph("正文在图片之前。")
            doc.add_picture(str(image_path))

        data = _build_docx(build)
        with pytest.raises(NeedsMultimodalParsingException):
            DocxParser().parse(_docx_input(data))


class TestMetadata:
    def test_metadata_includes_format_and_filename(self):
        def build(doc):
            doc.add_paragraph("有内容")

        data = _build_docx(build)
        result = DocxParser().parse(_docx_input(data, filename="需求说明.docx"))
        assert result.metadata["format"] == "docx"
        assert result.metadata["filename"] == "需求说明.docx"
