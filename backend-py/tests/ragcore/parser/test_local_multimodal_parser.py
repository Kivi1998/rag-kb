"""LocalMultimodalDocumentParser 测试。

deepdoc OCR 真实模型加载需要下载约 90MB 权重(见
`tests/ragcore/parser/deepdoc/test_deepdoc_pipeline.py` 里 `TestParsePageEndToEnd`
的 skip 说明),单测里一律 monkeypatch 掉 `_safe_parse_page`/`_ocr_bytes` 这几个
"真正调用模型"的边界函数,只验证:
- `supports()` 开关 × 扩展名矩阵;
- 各扩展名正确分发到对应处理逻辑(PDF 原生文本页 vs OCR 页、DOCX/PPTX 内嵌图片
  OCR、图片直接 OCR);
- OCR 产出空结果时的 `ParseError` 语义;
- 解析器返回值真的是 `await` 得到的结果(而不是协程对象本身),覆盖
  `DocumentParserFactory` 侧需要的"parse() 是 async def"这一契约。

真实 OCR 全链路已经在任务报告里用真实文件手动验证过一次,不在这里重复跑真实
模型推理。2026-09 已移除音视频(ASR)支持,原音视频落盘转写相关测试一并删除。
"""

from __future__ import annotations

import io

import pytest
from docx import Document
from PIL import Image

from app.ragcore.parser import local_multimodal_parser as lmp
from app.ragcore.parser.local_multimodal_parser import (
    LocalMultimodalDocumentParser,
    ParseError,
    ParserUnavailableError,
)
from tests.ragcore.parser.conftest import FakeDocumentInput


def _input(data: bytes, filename: str, content_type: str | None = None) -> FakeDocumentInput:
    return FakeDocumentInput(data, filename, content_type=content_type)


class TestSupports:
    def test_rejects_when_disabled(self, monkeypatch):
        monkeypatch.setattr(lmp.settings, "rag_multimodal_parser_enabled", False)
        parser = LocalMultimodalDocumentParser()
        assert parser.supports(None, "scan.pdf") is False

    def test_supports_all_documented_extensions_when_enabled(self, monkeypatch):
        monkeypatch.setattr(lmp.settings, "rag_multimodal_parser_enabled", True)
        parser = LocalMultimodalDocumentParser()
        for ext in ("pdf", "docx", "pptx", "png", "jpg", "jpeg", "tif", "tiff"):
            assert parser.supports(None, f"file.{ext}") is True, ext

    def test_rejects_unsupported_extension(self, monkeypatch):
        monkeypatch.setattr(lmp.settings, "rag_multimodal_parser_enabled", True)
        parser = LocalMultimodalDocumentParser()
        assert parser.supports(None, "notes.txt") is False
        assert parser.supports(None, "archive.zip") is False

    def test_rejects_when_no_filename(self, monkeypatch):
        monkeypatch.setattr(lmp.settings, "rag_multimodal_parser_enabled", True)
        parser = LocalMultimodalDocumentParser()
        assert parser.supports("application/pdf", None) is False


class TestParseDispatch:
    """验证 parse() 按扩展名分发到正确的内部处理函数(mock 掉真正的模型调用)。"""

    async def test_unsupported_extension_raises_parse_error(self):
        parser = LocalMultimodalDocumentParser()
        with pytest.raises(ParseError, match="不支持的扩展名"):
            await parser.parse(_input(b"whatever", "archive.zip"))

    async def test_pdf_native_text_page_skips_ocr(self, monkeypatch):
        import fitz

        called = {"ocr": False}

        def _fake_safe_parse_page(image):
            called["ocr"] = True
            return []

        monkeypatch.setattr(lmp, "_safe_parse_page", _fake_safe_parse_page)

        doc = fitz.open()
        page = doc.new_page()
        long_text = "这是一段足够长的原生文本内容。" * 20
        page.insert_text((72, 72), long_text)
        content = doc.tobytes()
        doc.close()

        parser = LocalMultimodalDocumentParser()
        result = await parser.parse(_input(content, "native.pdf"))

        assert called["ocr"] is False
        assert result.metadata["ocrPageCount"] == "0"
        assert result.metadata["parser"] == "local-multimodal"
        assert len(result.sections) == 1
        assert result.sections[0].metadata["ocr"] == "false"

    async def test_pdf_low_text_page_falls_back_to_ocr(self, monkeypatch):
        import fitz

        def _fake_safe_parse_page(image):
            return [{"text": "扫描页识别出的文字", "layout_type": "text"}]

        monkeypatch.setattr(lmp, "_safe_parse_page", _fake_safe_parse_page)

        doc = fitz.open()
        doc.new_page()  # 空白页,原生文本量为 0
        content = doc.tobytes()
        doc.close()

        parser = LocalMultimodalDocumentParser()
        result = await parser.parse(_input(content, "scan.pdf"))

        assert result.metadata["ocrPageCount"] == "1"
        assert len(result.sections) == 1
        assert result.sections[0].text == "扫描页识别出的文字"
        assert result.sections[0].metadata["ocr"] == "true"

    async def test_pdf_all_pages_empty_raises_parse_error(self, monkeypatch):
        import fitz

        monkeypatch.setattr(lmp, "_safe_parse_page", lambda image: [])

        doc = fitz.open()
        doc.new_page()
        content = doc.tobytes()
        doc.close()

        parser = LocalMultimodalDocumentParser()
        with pytest.raises(ParseError, match="PDF 未提取到任何可索引文本"):
            await parser.parse(_input(content, "blank.pdf"))

    async def test_docx_extracts_paragraph_and_ocrs_embedded_image(self, monkeypatch, tmp_path):
        monkeypatch.setattr(lmp, "_ocr_bytes", lambda data: "图片里识别出的文字")

        image_path = tmp_path / "figure.png"
        Image.new("RGB", (4, 4), color="red").save(image_path)

        document = Document()
        document.add_paragraph("这是正文段落。")
        document.add_picture(str(image_path))
        buf = io.BytesIO()
        document.save(buf)

        parser = LocalMultimodalDocumentParser()
        result = await parser.parse(_input(buf.getvalue(), "report.docx"))

        texts = [s.text for s in result.sections]
        assert any("这是正文段落" in t for t in texts)
        assert any("图片里识别出的文字" in t for t in texts)
        assert result.metadata["ocrImageCount"] == "1"

    async def test_docx_no_text_and_no_ocr_text_raises_parse_error(self, monkeypatch):
        monkeypatch.setattr(lmp, "_ocr_bytes", lambda data: "")

        document = Document()
        buf = io.BytesIO()
        document.save(buf)

        parser = LocalMultimodalDocumentParser()
        with pytest.raises(ParseError, match="DOCX 未提取到任何可索引文本"):
            await parser.parse(_input(buf.getvalue(), "empty.docx"))

    async def test_pptx_extracts_text_and_ocrs_picture_shape(self, monkeypatch, tmp_path):
        from pptx import Presentation

        monkeypatch.setattr(lmp, "_ocr_bytes", lambda data: "幻灯片图片里的文字")

        image_path = tmp_path / "slide_image.png"
        Image.new("RGB", (4, 4), color="blue").save(image_path)

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(0, 0, 100, 100)
        textbox.text_frame.text = "标题文字"
        slide.shapes.add_picture(str(image_path), 0, 0)

        buf = io.BytesIO()
        presentation.save(buf)

        parser = LocalMultimodalDocumentParser()
        result = await parser.parse(_input(buf.getvalue(), "deck.pptx"))

        assert len(result.sections) == 1
        assert "标题文字" in result.sections[0].text
        assert "幻灯片图片里的文字" in result.sections[0].text
        assert result.metadata["ocrImageCount"] == "1"

    async def test_pptx_no_text_anywhere_raises_parse_error(self, monkeypatch):
        from pptx import Presentation

        monkeypatch.setattr(lmp, "_ocr_bytes", lambda data: "")

        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        buf = io.BytesIO()
        presentation.save(buf)

        parser = LocalMultimodalDocumentParser()
        with pytest.raises(ParseError, match="PPTX 未提取到任何可索引文本"):
            await parser.parse(_input(buf.getvalue(), "empty.pptx"))

    async def test_image_extension_routes_to_ocr(self, monkeypatch):
        monkeypatch.setattr(
            lmp, "_safe_parse_page",
            lambda image: [{"text": "图片正文", "layout_type": "text"}],
        )

        img = Image.new("RGB", (10, 10), color="white")
        buf = io.BytesIO()
        img.save(buf, format="PNG")

        parser = LocalMultimodalDocumentParser()
        result = await parser.parse(_input(buf.getvalue(), "scan.png"))

        assert len(result.sections) == 1
        assert result.sections[0].text == "图片正文"
        assert result.metadata["format"] == "png"

    async def test_image_ocr_empty_result_raises_parse_error(self, monkeypatch):
        monkeypatch.setattr(lmp, "_safe_parse_page", lambda image: [])

        img = Image.new("RGB", (10, 10), color="white")
        buf = io.BytesIO()
        img.save(buf, format="PNG")

        parser = LocalMultimodalDocumentParser()
        with pytest.raises(ParseError, match="图片 OCR 未产出任何文字"):
            await parser.parse(_input(buf.getvalue(), "blank.png"))


class TestOcrModelUnavailable:
    def test_safe_parse_page_wraps_model_failure(self, monkeypatch):
        def _boom(image):
            raise RuntimeError("model weights missing")

        monkeypatch.setattr(lmp.deepdoc_pipeline, "parse_page", _boom)
        with pytest.raises(ParserUnavailableError, match="OCR/版面分析模型不可用"):
            lmp._safe_parse_page(Image.new("RGB", (2, 2)))

    def test_ocr_bytes_wraps_model_load_failure(self, monkeypatch):
        def _boom():
            raise RuntimeError("failed to download weights")

        monkeypatch.setattr(lmp.deepdoc_pipeline, "ensure_models_ready", _boom)
        with pytest.raises(ParserUnavailableError, match="OCR/版面分析模型不可用"):
            lmp._ocr_bytes(b"not-a-real-image")

    def test_ocr_bytes_skips_undecodable_single_image(self, monkeypatch):
        """单张图片解码失败不该让整份文档解析失败——按无文字处理，跳过。"""
        monkeypatch.setattr(lmp.deepdoc_pipeline, "ensure_models_ready", lambda: None)
        result = lmp._ocr_bytes(b"not-a-real-image-body")
        assert result == ""
