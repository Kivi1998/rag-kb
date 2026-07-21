"""PptxParser 测试。

用 `python-pptx` 现造最小 pptx 文档验证解析结果结构（不是逐字节对齐 Java 版）：
每页用第一个短文本框作为标题、page_no/slideNo 写入元数据、无文字的幻灯片被跳过、
全文档没有任何文字时触发 `NeedsMultimodalParsingException`。
"""

from __future__ import annotations

import io

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.ragcore.model import NeedsMultimodalParsingException
from app.ragcore.parser.pptx_parser import PptxParser
from tests.ragcore.parser.conftest import FakeDocumentInput

_PPTX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
_BLANK_LAYOUT_INDEX = 6  # python-pptx 默认模板里的 Blank 版式，没有占位符文字干扰断言


def _add_textbox(slide, text: str, top_inches: float = 1.0) -> None:
    box = slide.shapes.add_textbox(Inches(1), Inches(top_inches), Inches(6), Inches(1))
    box.text_frame.text = text


def _build_pptx(build_fn) -> bytes:
    presentation = Presentation()
    build_fn(presentation)
    buf = io.BytesIO()
    presentation.save(buf)
    return buf.getvalue()


def _pptx_input(data: bytes, filename: str = "deck.pptx") -> FakeDocumentInput:
    return FakeDocumentInput(data, filename, content_type=_PPTX_CONTENT_TYPE)


class TestSupports:
    def test_supports_by_content_type(self):
        assert PptxParser().supports(_PPTX_CONTENT_TYPE, None) is True

    def test_supports_by_extension(self):
        assert PptxParser().supports(None, "deck.pptx") is True

    def test_rejects_unrelated(self):
        assert PptxParser().supports(None, "deck.key") is False


class TestSlideTextExtraction:
    def test_first_short_textbox_becomes_title(self):
        def build(p):
            slide = p.slides.add_slide(p.slide_layouts[_BLANK_LAYOUT_INDEX])
            _add_textbox(slide, "第一页标题", top_inches=1.0)
            _add_textbox(slide, "这里是第一页正文内容。", top_inches=2.0)

        data = _build_pptx(build)
        result = PptxParser().parse(_pptx_input(data))

        assert len(result.sections) == 1
        section = result.sections[0]
        assert section.heading_path == ("第一页标题",)
        assert "第一页标题" in section.text
        assert "这里是第一页正文内容" in section.text
        assert section.page_no == 1
        assert section.metadata["slideNo"] == "1"
        assert section.metadata["type"] == "slide"

    def test_slide_without_short_line_falls_back_to_slide_number_title(self):
        long_text = "这是一段没有短文本框可以当标题的很长很长很长很长很长很长的文字内容超过一百二十个字符" * 3

        def build(p):
            slide = p.slides.add_slide(p.slide_layouts[_BLANK_LAYOUT_INDEX])
            _add_textbox(slide, long_text)

        data = _build_pptx(build)
        result = PptxParser().parse(_pptx_input(data))

        assert result.sections[0].heading_path == ("第 1 页",)

    def test_multiple_slides_get_incrementing_page_no(self):
        def build(p):
            for i in range(3):
                slide = p.slides.add_slide(p.slide_layouts[_BLANK_LAYOUT_INDEX])
                _add_textbox(slide, f"第{i + 1}页")

        data = _build_pptx(build)
        result = PptxParser().parse(_pptx_input(data))

        page_numbers = [s.page_no for s in result.sections]
        assert page_numbers == [1, 2, 3]

    def test_slide_with_no_text_is_skipped(self):
        def build(p):
            p.slides.add_slide(p.slide_layouts[_BLANK_LAYOUT_INDEX])  # 空白页，没有文字
            slide2 = p.slides.add_slide(p.slide_layouts[_BLANK_LAYOUT_INDEX])
            _add_textbox(slide2, "有内容的第二页")

        data = _build_pptx(build)
        result = PptxParser().parse(_pptx_input(data))

        assert len(result.sections) == 1
        assert result.sections[0].page_no == 2


class TestNoTextRaises:
    def test_all_slides_empty_raises_needs_multimodal_parsing(self):
        def build(p):
            p.slides.add_slide(p.slide_layouts[_BLANK_LAYOUT_INDEX])
            p.slides.add_slide(p.slide_layouts[_BLANK_LAYOUT_INDEX])

        data = _build_pptx(build)
        with pytest.raises(NeedsMultimodalParsingException):
            PptxParser().parse(_pptx_input(data))

    def test_no_slides_at_all_raises_needs_multimodal_parsing(self):
        data = _build_pptx(lambda p: None)
        with pytest.raises(NeedsMultimodalParsingException):
            PptxParser().parse(_pptx_input(data))


class TestMetadata:
    def test_metadata_includes_format_and_slide_count(self):
        def build(p):
            slide = p.slides.add_slide(p.slide_layouts[_BLANK_LAYOUT_INDEX])
            _add_textbox(slide, "内容")

        data = _build_pptx(build)
        result = PptxParser().parse(_pptx_input(data))
        assert result.metadata["format"] == "pptx"
        assert result.metadata["slideCount"] == "1"
