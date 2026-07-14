from __future__ import annotations

import io
import os
import threading
from pathlib import Path
from typing import Any


class ParseError(RuntimeError):
    pass


class ParserUnavailableError(RuntimeError):
    pass


def parse_document(path: Path, filename: str) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix == ".pptx":
        return _parse_pptx(path)
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
        return _parse_image(path)
    if suffix in {".mp3", ".wav", ".mp4", ".mov", ".mkv"}:
        return _parse_media(path)
    raise ParseError(f"unsupported format: {suffix}")


def _parse_pdf(path: Path) -> dict[str, Any]:
    import fitz
    from PIL import Image

    from .deepdoc.pipeline import parse_page as deepdoc_parse_page

    document = fitz.open(path)
    try:
        sections: list[dict[str, Any]] = []
        ocr_pages = 0
        for index, page in enumerate(document):
            text = page.get_text("text").strip()
            if len(text) >= 100:
                # 可提取文本页：直接用原生文本，速度快、无需模型。
                # 注：这条路径目前不做版面分析，多栏文本型 PDF 仍可能按行拼接；
                # 后续如需修复，思路是改用 page.get_text("words") 取带位置信息
                # 的文字框，复用下面 OCR 分支同样的 LayoutRecognizer 分组逻辑。
                sections.append(_section(text, page_no=index + 1,
                                         metadata={"type": "page", "ocr": "false"}))
                continue

            # 扫描页/低文本页：渲染成图片，走 deepdoc 的 OCR + 版面分析 + 表格结构
            # 识别流水线（替代原来的 RapidOCR 纯文字识别），解决多栏错拼和无边框
            # 表格错位的问题。
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image = Image.open(io.BytesIO(pixmap.tobytes("png")))
            page_sections = deepdoc_parse_page(image)
            ocr_pages += 1
            for sec in page_sections:
                sections.append(_section(sec["text"], page_no=index + 1,
                                         metadata={"type": sec.get("layout_type") or "page", "ocr": "true"}))
        if not sections:
            raise ParseError("PDF contains no extractable text")
        return {"sections": sections, "metadata": {
            "format": "pdf", "pageCount": str(len(document)), "ocrPageCount": str(ocr_pages)
        }}
    finally:
        document.close()


def _parse_docx(path: Path) -> dict[str, Any]:
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    document = Document(path)
    sections: list[dict[str, Any]] = []
    headings: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name.lower() if paragraph.style else ""
        if style.startswith("heading"):
            level_text = "".join(ch for ch in style if ch.isdigit())
            level = int(level_text or "1")
            headings = headings[: max(0, level - 1)] + [text]
        else:
            sections.append(_section(text, heading_path=headings, metadata={"type": "paragraph"}))
    for table_index, table in enumerate(document.tables, start=1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        text = "\n".join(row for row in rows if row.strip())
        if text:
            sections.append(_section(text, metadata={"type": "table", "tableIndex": str(table_index)}))
    ocr_images = 0
    seen_parts: set[str] = set()
    for relationship in document.part.rels.values():
        if relationship.reltype != RT.IMAGE:
            continue
        part_name = str(relationship.target_part.partname)
        if part_name in seen_parts:
            continue
        seen_parts.add(part_name)
        image_text = _ocr_bytes(relationship.target_part.blob)
        if image_text:
            ocr_images += 1
            sections.append(_section(image_text, metadata={"type": "image", "ocr": "true"}))
    if not sections:
        raise ParseError("DOCX contains no extractable text")
    return {"sections": sections, "metadata": {"format": "docx", "ocrImageCount": str(ocr_images)}}


def _parse_pptx(path: Path) -> dict[str, Any]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(path)
    sections: list[dict[str, Any]] = []
    ocr_images = 0
    for slide_no, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        title = ""
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    if not title and len(text) <= 120:
                        title = text.splitlines()[0]
                    lines.append(text)
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image_text = _ocr_bytes(shape.image.blob)
                if image_text:
                    lines.append(image_text)
                    ocr_images += 1
        if lines:
            sections.append(_section("\n".join(lines), heading_path=[title] if title else [],
                                     page_no=slide_no, metadata={"type": "slide"}))
    if not sections:
        raise ParseError("PPTX contains no extractable text")
    return {"sections": sections, "metadata": {
        "format": "pptx", "slideCount": str(len(presentation.slides)), "ocrImageCount": str(ocr_images)
    }}


def _parse_image(path: Path) -> dict[str, Any]:
    from PIL import Image

    from .deepdoc.pipeline import parse_page as deepdoc_parse_page

    image = Image.open(path)
    page_sections = deepdoc_parse_page(image)
    if not page_sections:
        raise ParseError("image OCR produced no text")
    sections = [_section(sec["text"], page_no=1,
                         metadata={"type": sec.get("layout_type") or "image", "ocr": "true"})
                for sec in page_sections]
    return {"sections": sections,
            "metadata": {"format": path.suffix.lower().lstrip("."), "ocrPageCount": "1"}}


_ocr_lock = threading.Lock()


def ensure_ocr_ready() -> str:
    """确保 deepdoc OCR 模型已加载（供 /ready 健康检查探测，也被
    _ocr_bytes 用于文档内嵌图片的纯文字识别，不做版面/表格分析）。"""
    from .deepdoc import pipeline
    with _ocr_lock:
        if pipeline._ocr is None:
            try:
                from .deepdoc.vision import OCR
                pipeline._ocr = OCR()
            except Exception as exc:
                raise ParserUnavailableError("OCR model is unavailable") from exc
    return "deepdoc-ocr"


def ocr_is_loaded() -> bool:
    from .deepdoc import pipeline
    return pipeline._ocr is not None


def _ocr_bytes(value: bytes) -> str:
    """给 DOCX/PPTX 里内嵌的单张图片（截图、Logo 等）做纯文字识别。
    这类图片一般不是多栏/表格版面，不需要跑版面分析和表格结构识别，
    直接用 OCR 模型识别文字即可，比跑完整 parse_page 更快。"""
    import numpy as np
    from PIL import Image

    ensure_ocr_ready()
    from .deepdoc import pipeline
    with _ocr_lock:
        np_image = np.array(Image.open(io.BytesIO(value)).convert("RGB"))
        raw_boxes = pipeline._ocr(np_image)
    if not raw_boxes:
        return ""
    return "\n".join(text.strip() for _, (text, _score) in raw_boxes if text and text.strip())


_whisper_model: Any = None
_whisper_lock = threading.Lock()


def ensure_asr_ready() -> str:
    global _whisper_model
    try:
        from faster_whisper import WhisperModel
    except ImportError as exc:
        raise ParserUnavailableError("ASR runtime is not installed") from exc
    with _whisper_lock:
        if _whisper_model is None:
            model_name = os.getenv("PARSER_WHISPER_MODEL", "small")
            device = os.getenv("PARSER_WHISPER_DEVICE", "cpu")
            compute_type = os.getenv("PARSER_WHISPER_COMPUTE_TYPE", "int8" if device == "cpu" else "float16")
            try:
                _whisper_model = WhisperModel(model_name, device=device, compute_type=compute_type)
            except Exception as exc:
                raise ParserUnavailableError(f"ASR model is unavailable: {model_name}") from exc
        return os.getenv("PARSER_WHISPER_MODEL", "small")


def asr_is_loaded() -> bool:
    return _whisper_model is not None


def _parse_media(path: Path) -> dict[str, Any]:
    ensure_asr_ready()
    with _whisper_lock:
        segments, info = _whisper_model.transcribe(str(path), vad_filter=True)
        sections = [_section(segment.text.strip(), metadata={
            "type": "subtitle", "timeRange": f"{_timestamp(segment.start)} --> {_timestamp(segment.end)}",
            "startSeconds": f"{segment.start:.3f}", "endSeconds": f"{segment.end:.3f}"
        }) for segment in segments if segment.text.strip()]
    if not sections:
        raise ParseError("ASR produced no transcript")
    return {"sections": sections, "metadata": {
        "format": path.suffix.lower().lstrip("."), "language": str(info.language), "segmentCount": str(len(sections))
    }}


def _section(text: str, heading_path: list[str] | None = None, page_no: int | None = None,
             metadata: dict[str, str] | None = None) -> dict[str, Any]:
    return {"text": text, "headingPath": heading_path or [], "pageNo": page_no,
            "lineStart": None, "lineEnd": None, "metadata": metadata or {}}


def _timestamp(seconds: float) -> str:
    total_ms = int(seconds * 1000)
    hours, remainder = divmod(total_ms, 3_600_000)
    minutes, remainder = divmod(remainder, 60_000)
    secs, millis = divmod(remainder, 1000)
    return f"{hours:02}:{minutes:02}:{secs:02}.{millis:03}"
