#!/usr/bin/env python3
"""生成可重复的脱敏多模态验收样本。

2026-09:已移除音视频（ASR）样本生成——当前业务不再需要音视频文档解析，原
espeak-ng/ffmpeg 生成的 wav/mp3/mp4/mov/mkv 五个 fixture 一并删除，对应权重
（100）转移给 markdown/code 两个既有类别，保持总数 5000 不变。
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import fitz
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches as DocxInches
from pptx import Presentation
from pptx.util import Inches


EXPECTED = "RAG PRODUCTION ACCEPTANCE 2026"


def make_image(path: Path) -> None:
    image = Image.new("RGB", (1800, 700), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 94)
    draw.text((80, 230), EXPECTED, fill="black", font=font)
    image.save(path)


def make_docx(path: Path) -> None:
    document = Document()
    if "Code Block" not in [style.name for style in document.styles]:
        document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    document.add_heading("Production Ingestion", level=1)
    document.add_paragraph(f"{EXPECTED}: Word paragraph with architecture and recovery guidance.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "component"
    table.cell(0, 1).text = "status"
    table.cell(1, 0).text = "parser"
    table.cell(1, 1).text = "ready"
    document.save(path)


def make_image_docx(path: Path, image_path: Path) -> None:
    document = Document()
    document.add_heading("Scanned Word Evidence", level=1)
    document.add_picture(str(image_path), width=DocxInches(6.5))
    document.save(path)


def make_pptx(path: Path, image_path: Path, image_only: bool) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    if image_only:
        slide.shapes.add_picture(str(image_path), Inches(0.5), Inches(1.2), width=Inches(12.0))
    else:
        title = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11), Inches(1))
        title.text_frame.text = "Production Ingestion"
        body = slide.shapes.add_textbox(Inches(0.8), Inches(2), Inches(11), Inches(3))
        body.text_frame.text = f"{EXPECTED}\nPPTX text, page and heading extraction."
    presentation.save(path)


def make_pdf(path: Path, image_path: Path, scanned: bool) -> None:
    document = fitz.open()
    page = document.new_page(width=900, height=500)
    if scanned:
        page.insert_image(page.rect, filename=str(image_path))
    else:
        text = (f"{EXPECTED} Text PDF acceptance. " * 8).strip()
        page.insert_textbox(fitz.Rect(50, 60, 850, 440), text, fontsize=18)
    document.save(path)
    document.close()


def main(output: Path) -> None:
    output.mkdir(parents=True, exist_ok=True)
    image = output / "ocr.png"
    make_image(image)
    Image.open(image).save(output / "ocr.jpg", quality=95)
    Image.open(image).save(output / "ocr.tiff")
    make_docx(output / "technical.docx")
    make_image_docx(output / "word-image-only.docx", image)
    make_pptx(output / "slides-text.pptx", image, False)
    make_pptx(output / "slides-image-only.pptx", image, True)
    make_pdf(output / "text.pdf", image, False)
    make_pdf(output / "scanned.pdf", image, True)

    (output / "guide.md").write_text(
        f"# Production Ingestion\n\n{EXPECTED}\n\n```java\nclass Acceptance {{}}\n```\n", encoding="utf-8")
    (output / "service.java").write_text(
        f"// {EXPECTED}\nclass AcceptanceService {{ static final String CODE = \"ERR_2026\"; }}\n", encoding="utf-8")
    (output / "config.json").write_text(
        json.dumps({"marker": EXPECTED, "worker": {"lease": True}}), encoding="utf-8")
    (output / "captions.srt").write_text(
        f"1\n00:00:00,000 --> 00:00:04,000\n{EXPECTED}\n", encoding="utf-8")
    (output / "captions.vtt").write_text(
        f"WEBVTT\n\n00:00:00.000 --> 00:00:04.000\n{EXPECTED}\n", encoding="utf-8")

    fixtures = [
        {"name": "guide.md", "kind": "markdown", "expected": EXPECTED, "weight": 1350},
        {"name": "service.java", "kind": "code", "expected": EXPECTED, "weight": 950},
        {"name": "config.json", "kind": "structured", "expected": EXPECTED, "weight": 500},
        {"name": "technical.docx", "kind": "docx", "expected": EXPECTED, "weight": 850},
        {"name": "word-image-only.docx", "kind": "docx-ocr", "expected": "RAG PRODUCTION", "weight": 50},
        {"name": "slides-text.pptx", "kind": "pptx-text", "expected": EXPECTED, "weight": 400},
        {"name": "text.pdf", "kind": "pdf-text", "expected": EXPECTED, "weight": 350},
        {"name": "scanned.pdf", "kind": "pdf-ocr", "expected": "RAG PRODUCTION", "weight": 100},
        {"name": "slides-image-only.pptx", "kind": "pptx-ocr", "expected": "RAG PRODUCTION", "weight": 100},
        {"name": "ocr.png", "kind": "image-png", "expected": "RAG PRODUCTION", "weight": 100},
        {"name": "ocr.jpg", "kind": "image-jpg", "expected": "RAG PRODUCTION", "weight": 100},
        {"name": "ocr.tiff", "kind": "image-tiff", "expected": "RAG PRODUCTION", "weight": 50},
        {"name": "captions.srt", "kind": "subtitle-srt", "expected": EXPECTED, "weight": 50},
        {"name": "captions.vtt", "kind": "subtitle-vtt", "expected": EXPECTED, "weight": 50},
    ]
    assert sum(item["weight"] for item in fixtures) == 5000
    (output / "fixtures.json").write_text(json.dumps(fixtures, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"fixtures": len(fixtures), "documents": 5000, "path": str(output)}))


if __name__ == "__main__":
    main(Path(sys.argv[1] if len(sys.argv) > 1 else "/fixtures"))
